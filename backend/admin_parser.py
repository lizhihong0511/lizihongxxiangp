"""Word/PDF 文档解析 + AI 提取题目"""
import os, json, tempfile
from openai import OpenAI

_client = None

EXTRACT_SYSTEM_PROMPT = """你是一个专业的考试命题专家。从以下文档内容中提取所有练习题。

对于每道题，判断其类型并提取：
- single_choice（单选题）：需提取选项 A/B/C/D
- multi_choice（多选题）：需提取选项 A/B/C/D
- true_false（判断题）：判断正误
- fill_blank（填空题）：空格处答案

返回严格 JSON（不要带 markdown 标记）：
{
  "questions": [
    {
      "type": "single_choice | multi_choice | true_false | fill_blank",
      "difficulty": "easy | medium | hard",
      "question_text": "题目内容",
      "options": [{"label": "A", "text": "选项内容"}],
      "correct_answer": "A | A,B | True/False | 答案文本",
      "explanation": "解析"
    }
  ]
}

要求：
1. 如文档中没有明确题目，根据文档内容自动生成题目
2. 干扰项要有迷惑性
3. 解析要包含知识点回顾
4. difficulty 根据题目难度合理判断"""


def _get_client() -> OpenAI:
    global _client
    if _client is None:
        _client = OpenAI(
            api_key=os.getenv("DEEPSEEK_API_KEY"),
            base_url="https://api.deepseek.com",
        )
    return _client


def extract_text_from_docx(filepath: str) -> str:
    """从 .docx 文件提取纯文本"""
    from docx import Document
    doc = Document(filepath)
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            lines.append(text)
    # 也提取表格内容
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def extract_text_from_pdf(filepath: str) -> str:
    """从 .pdf 文件提取纯文本"""
    import fitz  # PyMuPDF
    doc = fitz.open(filepath)
    lines = []
    for page in doc:
        text = page.get_text().strip()
        if text:
            lines.append(text)
    doc.close()
    return "\n".join(lines)


def extract_text(filepath: str) -> str:
    """自动识别文件类型并提取文本"""
    ext = os.path.splitext(filepath)[1].lower()
    if ext == ".docx":
        return extract_text_from_docx(filepath)
    elif ext == ".pdf":
        return extract_text_from_pdf(filepath)
    else:
        raise ValueError(f"不支持的文件格式: {ext}，仅支持 .docx 和 .pdf")


def extract_questions_from_text(text: str) -> list[dict]:
    """AI 从文本中提取题目"""
    client = _get_client()
    # 文本太长时截断（DeepSeek context window）
    max_chars = 60000
    if len(text) > max_chars:
        text = text[:max_chars] + "\n\n...(文档过长已截断)"

    response = client.chat.completions.create(
        model="deepseek-chat",
        messages=[
            {"role": "system", "content": EXTRACT_SYSTEM_PROMPT},
            {"role": "user", "content": f"以下是从文档中提取的内容，请从中提取题目：\n\n{text}"},
        ],
        response_format={"type": "json_object"},
        temperature=0.3,
        max_tokens=4096,
    )

    content = response.choices[0].message.content
    try:
        result = json.loads(content)
        questions = result.get("questions", [])
        # 验证每道题有必填字段
        valid = []
        for q in questions:
            if q.get("question_text") and q.get("correct_answer"):
                q.setdefault("options", [])
                q.setdefault("explanation", "")
                q.setdefault("difficulty", "medium")
                q.setdefault("type", "single_choice")
                valid.append(q)
        return valid
    except json.JSONDecodeError:
        raise ValueError("AI 解析文档内容失败，请重试")


def parse_and_extract(filepath: str) -> list[dict]:
    """完整流程：提取文本 → AI 提取题目"""
    text = extract_text(filepath)
    if not text.strip():
        raise ValueError("文档内容为空，无法提取题目")
    return extract_questions_from_text(text)
